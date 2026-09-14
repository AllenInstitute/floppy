#!/usr/bin/env python3
"""Keep resources/_text/ in step with resources/.

    refresh.py check    what changed since the last extraction (default)
    refresh.py run      re-extract everything stale
    refresh.py run --all  re-extract everything

The corpus is a live project folder. Answering from a stale extraction is worse
than being slow, so `check` is cheap enough to run whenever freshness matters.
"""
import argparse, hashlib, json, os, re, subprocess, sys, time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
RES = ROOT / "resources"
TEXT = RES / "_text"
STATE = TEXT / "MANIFEST.json"
EXTRACTABLE = {".docx", ".pdf", ".pptx", ".xlsx"}


def slug(rel: Path) -> str:
    return re.sub(r"[^A-Za-z0-9._ &-]", "_", str(rel).replace("/", "__"))


def sources():
    for p in sorted(RES.rglob("*")):
        if not p.is_file() or p.name == ".DS_Store":
            continue
        rel = p.relative_to(RES)
        if rel.parts[0] in ("_text", "schema"):
            continue
        if p.suffix.lower() in EXTRACTABLE:
            yield p, rel


def extract(p: Path) -> str:
    ext = p.suffix.lower()
    if ext == ".docx":
        return subprocess.run(["pandoc", "-t", "markdown", str(p)],
                              capture_output=True, text=True, timeout=300).stdout
    if ext == ".pdf":
        return subprocess.run(["pdftotext", "-layout", str(p), "-"],
                              capture_output=True, text=True, timeout=600).stdout
    if ext == ".xlsx":
        import openpyxl
        out = []
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets:
            out.append(f"\n### sheet: {ws.title}\n")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if cells:
                    out.append(" | ".join(cells))
        return "\n".join(out)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(p).slides, 1):
            blocks = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    blocks.append(sh.text_frame.text.strip())
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        blocks.append(" | ".join(c.text.strip() for c in r.cells))
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
                blocks.append("NOTES: " + s.notes_slide.notes_text_frame.text.strip())
            if blocks:
                out.append(f"\n## Slide {i}\n" + "\n\n".join(blocks))
        return "\n".join(out)
    return ""


def load_state():
    if not STATE.exists():
        return {}
    raw = json.loads(STATE.read_text())
    if isinstance(raw, list):
        raw = {r["source"]: r for r in raw}
    return {k: v for k, v in raw.items() if v.get("text")}


def survey():
    """Compare each source against the mtime recorded when it was last extracted."""
    state = load_state()
    new, changed, ok, gone = [], [], [], []
    seen = set()
    for p, rel in sources():
        seen.add(str(rel))
        rec = state.get(str(rel))
        if rec is None or not (TEXT / (slug(rel) + ".txt")).exists():
            new.append(rel)
        elif abs(p.stat().st_mtime - float(rec.get("mtime", -1))) > 1:
            changed.append(rel)
        else:
            ok.append(rel)
    for k, rec in state.items():
        # text-only entries have no bundled original to compare against
        if k not in seen and rec.get("original_bundled") is not False:
            gone.append(Path(k))
    return new, changed, ok, gone


def refresh_gitrefs():
    """Re-record the commit each bundled schema clone is sitting on.

    The gap-analysis house style requires citing a pinned version, so a GITREF
    that has drifted from the checkout is worse than none.
    """
    base = RES / "schema"
    if not base.exists():
        print("no resources/schema/ directory")
        return 1
    found = 0
    for repo in sorted(p for p in base.iterdir() if (p / ".git").exists()):
        r = subprocess.run(["git", "-C", str(repo), "log", "-1", "--format=%H %cd %s"],
                           capture_output=True, text=True)
        tag = subprocess.run(["git", "-C", str(repo), "describe", "--tags"],
                             capture_output=True, text=True).stdout.strip()
        line = r.stdout.strip() + (f"\ntag: {tag}" if tag else "")
        (base / f"{repo.name}.GITREF.txt").write_text(line + "\n")
        print(f"  {repo.name}: {line.splitlines()[0]}")
        found += 1
    stale = [p for p in base.iterdir()
             if p.is_dir() and not (p / ".git").exists() and not p.name.startswith(".")]
    for p in stale:
        print(f"  {p.name}: no .git — GITREF left as-is (snapshot, not a clone)")
    print(f"{found} clone(s) re-recorded.")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", nargs="?", default="check", choices=["check", "run", "schema"])
    ap.add_argument("--all", action="store_true")
    a = ap.parse_args()

    if a.cmd == "schema":
        return refresh_gitrefs()

    new, changed, ok, orphaned = survey()
    print(f"{len(ok)} up to date · {len(new)} new · {len(changed)} changed · {len(orphaned)} removed")
    for label, items in (("new", new), ("changed", changed), ("removed", orphaned)):
        for r in items:
            print(f"  {label:8} {r}")

    if a.cmd == "check":
        if new or changed:
            print("\nrun `python3 scripts/refresh.py run` to re-extract.")
        return 1 if (new or changed) else 0

    todo = [(p, rel) for p, rel in sources()
            if a.all or rel in new or rel in changed]
    if not todo:
        print("nothing to do")
        return 0
    TEXT.mkdir(parents=True, exist_ok=True)
    state = load_state()
    for p, rel in todo:
        try:
            txt = extract(p)
        except Exception as e:
            print(f"  FAILED {rel}: {e}", file=sys.stderr)
            continue
        if not txt.strip():
            print(f"  empty  {rel}")
            continue
        tpath = TEXT / (slug(rel) + ".txt")
        tpath.write_text(f"# source: {rel}\n# extracted: {time.strftime('%Y-%m-%d')}\n\n" + txt)
        state[str(rel)] = {"source": str(rel), "bytes": p.stat().st_size,
                           "mtime": p.stat().st_mtime, "text": tpath.name,
                           "original_bundled": True,
                           "extracted": time.strftime("%Y-%m-%d")}
        print(f"  wrote  {tpath.name}")
    STATE.write_text(json.dumps(state, indent=1, sort_keys=True))
    print(f"\n{len(todo)} file(s) re-extracted.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
