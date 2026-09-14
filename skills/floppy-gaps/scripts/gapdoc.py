#!/usr/bin/env python3
"""Build a BioData Registry gap analysis in the compact M0-Report format.

The JSON is the source of truth; the .docx is a rendering. Metrics are computed
from the field rows so the counts can never drift from the tables.

    gapdoc.py schema                        -> the expected JSON shape
    gapdoc.py seed metadata.nd.json -o a.json   -> skeleton JSON from a metadata export
    gapdoc.py validate  analysis.json       -> coherence checks (exit 1 if problems)
    gapdoc.py build     analysis.json       -> analysis.docx  (compact section)
    gapdoc.py metrics   analysis.json       -> print the counts, write nothing

The output matches the per-dataset sections of the Milestone 0 Report:

    <Dataset> Gap Analysis                       (title)
    Detailed Dataset Analysis: <descriptor>      (heading)
    <intro prose, optional labeled notes>
    Source documents reviewed: <bullets>
    Sample asset: <bullets>
    Gap Summary Overview        -> 3-col table  (Gap | Schema Layer | Detail)
    Field-Level Mapping         -> 6-col tables per entity
        (Source Field | Type | Target (aind) | Classification | Cross-Dataset Gap | Notes & Recommendation)

No numbered G/R/Q, no metrics table, no origin column — that is the older
10-section standalone format. This one is deliberately the short form: what the
schema can and cannot hold, and why, and nothing else.
"""
import argparse, json, sys
from collections import Counter
from pathlib import Path

CLASSES = ["Mapped", "Gap", "Conflict", "Unmappable", "Needs Review"]
FIELD_HEADERS = ["Source Field", "Type", "Target (aind)", "Classification",
                 "Cross-Dataset Gap", "Notes & Recommendation"]
ESCALATION = "Jess Thomas (jess.thomas@alleninstitute.org)"

PLACEHOLDER_BANNER = (
    "Placeholder note.",
    "Prepared from indirect material (a publication or a description), not from a "
    "metadata export. Source field names are descriptive labels for quantities the "
    "material reports, not real keys from a source system. Classifications are "
    "verified statements about the reference schema; the metrics are not comparable "
    "to an analysis built from a real field inventory.")

SCHEMA = {
    "mode": "full | placeholder  — placeholder stamps a note and marks metrics non-comparable",
    "meta": {
        "dataset": "str — dataset name, e.g. 'DNA Typewriter lineage recording'",
        "descriptor": "str — parenthetical for the heading, e.g. 'Mouse, SeaHub'",
        "source_organization": "str",
        "prepared_by": "str",
        "reference_schema": "str — MUST be pinned, e.g. 'aind-data-schema v2.9.0 (commit 18114d7b)'",
        "milestone": "str",
        "status": "Draft | In Review | Final",
        "date": "YYYY-MM-DD",
        "version": "str, e.g. v0.1",
    },
    "intro": ["str prose paragraph",
              {"label": "Placeholder note.", "body": "str — a labeled standalone paragraph"}],
    "sources_reviewed": ["str", "..."],
    "sample_asset": ["str bullet", "..."],
    "gap_summary": [{
        "gap": "str — short title; reuse the framing from known-gaps.md when it recurs",
        "layer": "str — e.g. 'Acquisition / DataDescription'",
        "detail": "str — 1–3 sentences; name the recurrence ('fifth appearance…') to reference prior work",
    }],
    "field_intro": "str — one sentence introducing the field tables (optional)",
    "entities": [{
        "name": "str — table title, e.g. 'Subject — transgenic embryo'",
        "fields": [{
            "source_field": "str",
            "source_type": "str",
            "target_field": "str — real aind path, or '— (no analog)' / '— (platform)'",
            "classification": "one of " + " | ".join(CLASSES),
            "cross_dataset_gap": "str — the cross-cutting gap this belongs to, or '—'",
            "notes": "str — mismatch detail and recommendation, combined",
        }],
    }],
    "post_gap_notes": [{"label": "Required-field readiness.", "body": "str — optional labeled paragraph(s)"}],
}


def all_rows(doc):
    return [f for e in doc.get("entities", []) for f in e.get("fields", [])]


def compute_metrics(doc):
    rows = all_rows(doc)
    counts = Counter(f.get("classification", "?") for f in rows)
    total = len(rows)
    pct = lambda n: f"{n} ({round(100 * n / total)}%)" if total else "0 (0%)"
    m = {"Total fields examined": str(total)}
    for c in CLASSES:
        m[c] = pct(counts.get(c, 0))
    m["Entities covered"] = str(len(doc.get("entities", [])))
    m["Gaps raised"] = str(len(doc.get("gap_summary", [])))
    off = {k: v for k, v in counts.items() if k not in CLASSES}
    return m, off


def validate(doc):
    problems = []
    meta = doc.get("meta", {})
    for k in ("dataset", "source_organization", "prepared_by", "reference_schema",
              "milestone", "status", "date", "version"):
        if not meta.get(k):
            problems.append(f"meta.{k} is missing or empty")
    ref = str(meta.get("reference_schema", ""))
    if "latest" in ref.lower() or not any(ch.isdigit() for ch in ref):
        problems.append(f"meta.reference_schema = {ref!r} is not pinned — cite a version "
                        "and ideally a commit; migration must not rest on 'latest'.")

    if doc.get("mode") not in ("full", "placeholder"):
        problems.append(f"mode = {doc.get('mode')!r}; use 'full' or 'placeholder'")

    if not doc.get("gap_summary"):
        problems.append("gap_summary is empty — it is the Gap Summary Overview table")
    for g in doc.get("gap_summary", []):
        for k in ("gap", "layer", "detail"):
            if not str(g.get(k, "")).strip():
                problems.append(f"a gap_summary row is missing {k!r}")

    rows = all_rows(doc)
    if not rows:
        problems.append("no field rows — the field tables are the substance of the document")
    _, off = compute_metrics(doc)
    for k, v in off.items():
        problems.append(f"{v} field row(s) use classification {k!r}, not one of {CLASSES}")
    for e in doc.get("entities", []):
        if not e.get("name"):
            problems.append("an entity has no name (used as the table title)")
        for f in e.get("fields", []):
            for k in ("source_field", "target_field", "classification"):
                if not str(f.get(k, "")).strip():
                    problems.append(f"a field in {e.get('name','?')!r} is missing {k!r}")

    for u in doc.get("unknowns", []):
        if not str(u.get("ask", "")).strip():
            problems.append(f"unknown {u.get('item')!r} names nobody to ask; default {ESCALATION}")
    return problems


# --------------------------------------------------------------- docx render

def build(doc, out):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.8)
    d.styles["Normal"].font.size = Pt(10)
    GREY = RGBColor(0x66, 0x66, 0x66)
    meta = doc["meta"]

    def para(text=""):
        return d.add_paragraph(text)

    def labeled(label, body):
        p = d.add_paragraph()
        p.add_run(label + " ").bold = True
        p.add_run(body)
        return p

    def bullets(items):
        for it in items:
            d.add_paragraph(str(it), style="List Bullet")

    def grid(headers, rows, widths, title=None):
        ncols = len(headers)
        t = d.add_table(rows=1, cols=ncols)
        t.style = "Table Grid"
        if title is not None:                     # merged title row on top
            trow = t.rows[0]
            merged = trow.cells[0].merge(trow.cells[ncols - 1])
            merged.text = ""
            r = merged.paragraphs[0].add_run(title); r.bold = True; r.font.size = Pt(9.5)
            hrow = t.add_row()
        else:
            hrow = t.rows[0]
        for i, h in enumerate(headers):
            hrow.cells[i].text = ""
            r = hrow.cells[i].paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9)
        for rec in rows:
            cells = t.add_row().cells
            for i, v in enumerate(rec):
                cells[i].text = ""
                cells[i].paragraphs[0].add_run("" if v is None else str(v)).font.size = Pt(9)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
        d.add_paragraph()
        return t

    # ---- title + provenance -------------------------------------------------
    title = doc.get("title") or f"{meta['dataset']} Gap Analysis"
    d.add_heading(title, 0)
    prov = d.add_paragraph()
    r = prov.add_run(f"{meta['source_organization']}  ·  {meta['prepared_by']}  ·  "
                     f"{meta['reference_schema']}  ·  {meta['milestone']}  ·  "
                     f"{meta['status']} {meta['version']}, {meta['date']}")
    r.font.size = Pt(8); r.font.color.rgb = GREY

    heading = doc.get("detailed_heading")
    if not heading:
        desc = meta.get("descriptor", "")
        heading = f"Detailed Dataset Analysis: {meta['dataset']}" + (f" ({desc})" if desc else "")
    d.add_heading(heading, 1)

    # ---- intro (strings or labeled objects) --------------------------------
    for item in doc.get("intro", []):
        if isinstance(item, dict):
            labeled(item.get("label", ""), item.get("body", ""))
        else:
            para(item)
    if doc.get("mode") == "placeholder" and not any(
            isinstance(x, dict) and x.get("label", "").startswith("Placeholder")
            for x in doc.get("intro", [])):
        labeled(*PLACEHOLDER_BANNER)

    if doc.get("sources_reviewed"):
        para("Source documents reviewed:")
        bullets(doc["sources_reviewed"])
    if doc.get("sample_asset"):
        para(doc.get("sample_lead", "Sample asset:"))
        bullets(doc["sample_asset"])

    # ---- Gap Summary Overview ----------------------------------------------
    d.add_heading("Gap Summary Overview", 2)
    grid(["Gap", "Schema Layer", "Detail"],
         [[g.get("gap"), g.get("layer"), g.get("detail")] for g in doc.get("gap_summary", [])],
         widths=[1.9, 1.3, 4.0])
    for note in doc.get("post_gap_notes", []):
        labeled(note.get("label", ""), note.get("body", ""))
    if doc.get("post_gap_notes"):
        d.add_paragraph()

    # ---- Field-Level Mapping -----------------------------------------------
    d.add_heading("Field-Level Mapping", 2)
    if doc.get("field_intro"):
        para(doc["field_intro"])
        d.add_paragraph()
    for e in doc.get("entities", []):
        grid(FIELD_HEADERS,
             [[f.get("source_field"), f.get("source_type"), f.get("target_field"),
               f.get("classification"), f.get("cross_dataset_gap", "—"), f.get("notes")]
              for f in e.get("fields", [])],
             widths=[1.1, 0.8, 1.2, 0.9, 1.1, 2.1],
             title=e.get("name"))

    # ---- footer: verification caveat ---------------------------------------
    p = d.add_paragraph()
    p.add_run(
        f"Reference schema: {meta['reference_schema']}. A mapping recorded here is a statement "
        "about schema compatibility, not a claim that the source field is populated or valid; "
        "re-verify against the pinned schema before the gate review."
    ).italic = True

    d.save(out)
    return compute_metrics(doc)[0]


# --------------------------------------------------------------- seed

_SKIP = {"describedBy", "schema_version", "$schema", "object_type"}


def _flatten(obj, prefix=""):
    """Yield (path, type_label) for the leaf/near-leaf fields of a metadata object."""
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k in _SKIP or k.startswith("_"):   # skip annotations & _COMMENT_ keys
                continue
            path = f"{prefix}.{k}" if prefix else k
            if isinstance(v, dict):
                yield from _flatten(v, path)
            elif isinstance(v, list):
                inner = "object[]" if v and isinstance(v[0], (dict, list)) else "array"
                yield path, inner
                if v and isinstance(v[0], dict):
                    yield from _flatten(v[0], path + "[]")
            else:
                yield path, type(v).__name__ if v is not None else "null"


CORE_TO_ENTITY = {
    "subject": "Subject",
    "data_description": "Data Description",
    "procedures": "Procedures",
    "acquisition": "Acquisition",
    "instrument": "Instrument",
    "processing": "Processing",
    "quality_control": "Quality Control",
    "metadata": "Metadata",
}


def _top_level_objects(text):
    """Yield every top-level {...} block, brace-matching and ignoring braces in
    strings. Tolerates hand-annotated exports where the outer object closes early
    and later core objects (instrument, processing) dangle as siblings."""
    depth = 0; start = None; in_str = False; esc = False
    for i, ch in enumerate(text):
        if in_str:
            if esc: esc = False
            elif ch == "\\": esc = True
            elif ch == '"': in_str = False
            continue
        if ch == '"': in_str = True
        elif ch == "{":
            if depth == 0: start = i
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0 and start is not None:
                yield text[start:i + 1]; start = None


def _load_metadata_objects(meta_path):
    text = Path(meta_path).read_text()
    try:
        return [json.loads(text)]
    except json.JSONDecodeError:
        pass
    objs = []
    for block in _top_level_objects(text):
        try:
            objs.append(json.loads(block))
        except json.JSONDecodeError:
            continue
    if not objs:
        raise SystemExit(f"could not parse any JSON object from {meta_path}")
    return objs


_OBJTYPE_TO_CORE = {
    "Subject": "subject", "DataDescription": "data_description",
    "Procedures": "procedures", "Acquisition": "acquisition",
    "Instrument": "instrument", "Processing": "processing",
    "QualityControl": "quality_control", "Metadata": "metadata",
}


def seed(meta_path):
    """Build a skeleton gap-analysis JSON from an aind-style metadata export.

    Walks each core object into a field inventory. source_field/source_type are
    filled from the export; target_field/classification/cross_dataset_gap/notes
    are left for the analyst to complete against the pinned schema. If the export
    is itself aind metadata (e.g. teaseq-metadata.nd.json), the source path is
    also the target path — the analyst confirms and reclassifies.
    """
    cores = {}
    for obj in _load_metadata_objects(meta_path):
        if not isinstance(obj, dict):
            continue
        # a wrapper carrying core keys directly
        hit = False
        for k, v in obj.items():
            if k in CORE_TO_ENTITY and isinstance(v, dict):
                cores.setdefault(k, v); hit = True
        if hit:
            continue
        # a standalone core object identified by object_type
        core = _OBJTYPE_TO_CORE.get(obj.get("object_type", ""))
        if core and core != "metadata":
            cores.setdefault(core, obj)
    if not cores:                       # flat single-object export
        cores = {"metadata": _load_metadata_objects(meta_path)[0]}

    entities = []
    for core, obj in cores.items():
        fields = []
        for path, typ in _flatten(obj):
            fields.append({
                "source_field": path,
                "source_type": typ,
                "target_field": "",
                "classification": "Needs Review",
                "cross_dataset_gap": "",
                "notes": "",
            })
        if fields:
            entities.append({"name": CORE_TO_ENTITY.get(core, core), "fields": fields})

    return {
        "mode": "full",
        "meta": {
            "dataset": Path(meta_path).stem,
            "descriptor": "",
            "source_organization": "",
            "prepared_by": "",
            "reference_schema": "aind-data-schema v2.9.0 (commit 18114d7b)",
            "milestone": "",
            "status": "Draft",
            "date": "",
            "version": "v0.1",
        },
        "intro": [f"Seeded from {Path(meta_path).name}. Replace this paragraph, complete the "
                  "empty target_field / classification / notes columns against the pinned "
                  "schema, and fill the gap_summary."],
        "sources_reviewed": [Path(meta_path).name,
                             "aind-data-schema v2.9.0, commit 18114d7b"],
        "sample_asset": [],
        "gap_summary": [],
        "field_intro": "These tables document field-level mappings for the core schemas.",
        "entities": entities,
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["schema", "seed", "validate", "build", "metrics"])
    ap.add_argument("json_file", nargs="?")
    ap.add_argument("-o", "--out", help="output path")
    a = ap.parse_args()

    if a.cmd == "schema":
        print(json.dumps(SCHEMA, indent=2)); return 0

    if not a.json_file:
        ap.error("json_file is required")

    if a.cmd == "seed":
        skeleton = seed(a.json_file)
        out = a.out or str(Path(a.json_file).with_suffix("")) + "-gap-seed.json"
        Path(out).write_text(json.dumps(skeleton, indent=1))
        n = sum(len(e["fields"]) for e in skeleton["entities"])
        print(f"wrote {out}  ({len(skeleton['entities'])} entities, {n} fields, "
              "classification=Needs Review — complete against the schema)")
        return 0

    doc = json.loads(Path(a.json_file).read_text())
    problems = validate(doc)
    if problems:
        print(f"{len(problems)} problem(s):", file=sys.stderr)
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
    else:
        print("validate: clean")

    if a.cmd == "validate":
        return 1 if problems else 0
    if a.cmd == "metrics":
        for k, v in compute_metrics(doc)[0].items():
            print(f"  {k}: {v}")
        return 1 if problems else 0

    out = a.out or str(Path(a.json_file).with_suffix("")) + ".docx"
    metrics = build(doc, out)
    print(f"wrote {out}")
    for k, v in metrics.items():
        print(f"  {k}: {v}")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
